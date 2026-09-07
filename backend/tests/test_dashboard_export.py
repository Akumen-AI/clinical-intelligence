import csv
import io
import json
import uuid

import pytest
from openpyxl import load_workbook

from app.core.security import create_access_token
from app.main import app
from app.models.audit_log import AuditLogEntry
from tests.conftest import TestingSessionLocal
from fastapi.testclient import TestClient


EXPORT_PAYLOAD = {
    "hospital": "main",
    "filters": {"department": None, "start_date": "2026-02-01", "end_date": "2026-02-28"},
    "metrics": [{
        "key": "admissions",
        "label": "Admissions",
        "value": 7,
        "unit": "visits",
        "chart": [{"date": "2026-02-01", "value": 1}],
        "available": True,
    }],
}


@pytest.mark.parametrize("export_format", ["pdf", "csv", "xlsx"])
def test_dashboard_export_contains_metrics_and_audit_log(monkeypatch, export_format):
    user_id = uuid.uuid4()
    token = create_access_token({"sub": str(user_id), "role": "hospital_admin", "email": "admin@example.com"})
    client = TestClient(app)
    client.headers.update({"Authorization": f"Bearer {token}"})
    monkeypatch.setattr("app.api.dashboards.get_hospital_dashboard", lambda *args, **kwargs: EXPORT_PAYLOAD)

    response = client.get(
        "/api/v1/dashboards/hospital/export",
        params={"format": export_format, "start_date": "2026-02-01", "end_date": "2026-02-28"},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"] == f'attachment; filename="operations-dashboard.{export_format}"'
    if export_format == "pdf":
        assert response.content.startswith(b"%PDF")
    elif export_format == "csv":
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert ["Admissions", "7", "visits", "Available"] == rows[6][:4]
    else:
        sheet = load_workbook(io.BytesIO(response.content), read_only=True).active
        assert list(sheet.iter_rows(min_row=7, max_row=7, values_only=True))[0][:4] == ("Admissions", "7", "visits", "Available")

    db = TestingSessionLocal()
    try:
        log = db.query(AuditLogEntry).filter(AuditLogEntry.actor_user_id == user_id).one()
        assert log.action_type == "dashboard_export"
        assert log.target_entity == "dashboard:hospital"
        assert json.loads(log.rationale)["format"] == export_format
    finally:
        db.close()


def test_dashboard_export_denies_unauthorized_role():
    token = create_access_token({"sub": str(uuid.uuid4()), "role": "doctor", "email": "doctor@example.com"})
    client = TestClient(app)
    client.headers.update({"Authorization": f"Bearer {token}"})

    response = client.get("/api/v1/dashboards/hospital/export", params={"format": "csv"})

    assert response.status_code == 403